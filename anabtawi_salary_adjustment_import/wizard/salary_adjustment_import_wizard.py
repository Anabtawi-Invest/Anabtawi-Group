import base64
import io
from datetime import date, datetime

from odoo import api, fields, models, _
from odoo.exceptions import UserError

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils.datetime import from_excel
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    openpyxl = None

TEMPLATE_FILENAME = "salary_adjustment_template.xlsx"
TEMPLATE_MAX_ROWS = 2000

# (key, header label, required header)
COLUMNS = [
    ("employee", "Employee", True),
    ("input_type", "Type", True),
    ("monthly_amount", "Payslip Amount", True),
    ("duration", "Duration", False),
    ("total_amount", "Total Amount", False),
    ("date_start", "Start Date", True),
    ("note", "Note", False),
]

DURATION_LABELS = {
    "one": "One Time",
    "limited": "Limited",
    "unlimited": "Unlimited",
}

DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%Y")


def _normalize(value):
    if value is None:
        return ""
    return " ".join(str(value).split()).casefold()


class SalaryAdjustmentImportWizard(models.TransientModel):
    _name = "salary.adjustment.import.wizard"
    _description = "Salary Adjustments Excel Import"

    state = fields.Selection(
        [("upload", "Upload"), ("preview", "Preview")],
        default="upload",
        required=True,
    )
    template_file = fields.Binary(readonly=True, attachment=False)
    template_name = fields.Char(default=TEMPLATE_FILENAME)
    excel_file = fields.Binary(string="Excel File", attachment=False)
    file_name = fields.Char()
    line_ids = fields.One2many("salary.adjustment.import.line", "wizard_id", string="Rows")
    valid_count = fields.Integer(compute="_compute_counts")
    error_count = fields.Integer(compute="_compute_counts")

    @api.depends("line_ids.status")
    def _compute_counts(self):
        for wizard in self:
            wizard.error_count = len(wizard.line_ids.filtered(lambda l: l.status == "error"))
            wizard.valid_count = len(wizard.line_ids) - wizard.error_count

    @api.onchange("excel_file")
    def _onchange_excel_file(self):
        self.state = "upload"
        self.line_ids = [fields.Command.clear()]

    def _reopen(self):
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
            "name": _("Import Salary Adjustments"),
        }

    def _check_openpyxl(self):
        if not openpyxl:
            raise UserError(_("The Python package 'openpyxl' is required to read and write Excel files."))

    def _get_input_types(self):
        return self.env["hr.payslip.input.type"].search([("available_in_attachments", "=", True)])

    # ------------------------------------------------------------------
    # Template
    # ------------------------------------------------------------------

    def action_download_template(self):
        self.ensure_one()
        self._check_openpyxl()

        input_types = self._get_input_types()
        employees = self.env["hr.employee"].search([("company_id", "in", self.env.companies.ids)], order="name")

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Salary Adjustments"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="714B67")
        widths = [40, 30, 18, 15, 18, 15, 40]
        for col_index, (_key, label, required) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=1, column=col_index, value=label + (" *" if required else ""))
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            sheet.column_dimensions[cell.column_letter].width = widths[col_index - 1]
        sheet.freeze_panes = "A2"

        for row in range(2, TEMPLATE_MAX_ROWS + 1):
            sheet.cell(row=row, column=3).number_format = "0.000"
            sheet.cell(row=row, column=5).number_format = "0.000"
            sheet.cell(row=row, column=6).number_format = "yyyy-mm-dd"

        lists = workbook.create_sheet("Lists")
        lists["A1"], lists["B1"], lists["C1"] = "Types", "Durations", "Employees"
        for index, input_type in enumerate(input_types, start=2):
            lists.cell(row=index, column=1, value=input_type.name)
        for index, label in enumerate(DURATION_LABELS.values(), start=2):
            lists.cell(row=index, column=2, value=label)
        for index, employee in enumerate(employees, start=2):
            lists.cell(row=index, column=3, value=employee.name)
        lists.sheet_state = "hidden"

        def add_list_validation(list_column, count, target_column, strict):
            if not count:
                return
            validation = DataValidation(
                type="list",
                formula1=f"=Lists!${list_column}$2:${list_column}${count + 1}",
                allow_blank=True,
                showErrorMessage=True,
                errorStyle="stop" if strict else "warning",
            )
            sheet.add_data_validation(validation)
            validation.add(f"{target_column}2:{target_column}{TEMPLATE_MAX_ROWS}")

        add_list_validation("C", len(employees), "A", strict=False)
        add_list_validation("A", len(input_types), "B", strict=True)
        add_list_validation("B", len(DURATION_LABELS), "D", strict=True)

        help_sheet = workbook.create_sheet("Instructions")
        help_sheet.column_dimensions["A"].width = 20
        help_sheet.column_dimensions["B"].width = 90
        instructions = [
            ("Column", "Description"),
            ("Employee *", "Employee name exactly as in Odoo (pick from the list)."),
            ("Type *", "Salary adjustment type (pick from the list)."),
            ("Payslip Amount *", "Amount applied on each payslip, must be greater than 0."),
            ("Duration", "One Time, Limited or Unlimited. Empty = One Time."),
            ("Total Amount", "Required only when Duration is Limited (must be >= Payslip Amount)."),
            ("Start Date *", "Date from which the adjustment applies (YYYY-MM-DD)."),
            ("Note", "Optional reason or reference."),
        ]
        for row_index, (col_a, col_b) in enumerate(instructions, start=1):
            help_sheet.cell(row=row_index, column=1, value=col_a)
            help_sheet.cell(row=row_index, column=2, value=col_b)
            if row_index == 1:
                help_sheet.cell(row=row_index, column=1).font = Font(bold=True)
                help_sheet.cell(row=row_index, column=2).font = Font(bold=True)

        buffer = io.BytesIO()
        workbook.save(buffer)
        self.write({
            "template_file": base64.b64encode(buffer.getvalue()),
            "template_name": TEMPLATE_FILENAME,
        })
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{self._name}/{self.id}/template_file/{TEMPLATE_FILENAME}?download=true",
            "target": "self",
        }

    # ------------------------------------------------------------------
    # Parsing
    # ------------------------------------------------------------------

    def _read_rows(self):
        """Return a list of (excel_row_number, {column_key: raw_value})."""
        self._check_openpyxl()
        if not self.excel_file:
            raise UserError(_("Please upload an Excel file first."))
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(base64.b64decode(self.excel_file)), data_only=True, read_only=True,
            )
        except Exception as error:
            raise UserError(_("The uploaded file is not a valid Excel (.xlsx) file.\n%s", error))

        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise UserError(_("The Excel file is empty."))

        labels = {_normalize(label): key for key, label, _required in COLUMNS}
        column_map = {}
        for index, title in enumerate(header):
            key = labels.get(_normalize(title).rstrip(" *").strip())
            if key and key not in column_map:
                column_map[key] = index

        missing = [label for key, label, required in COLUMNS if required and key not in column_map]
        if missing:
            raise UserError(_(
                "The following required columns are missing: %s.\nPlease use the downloaded template.",
                ", ".join(missing),
            ))

        result = []
        for row_number, values in enumerate(rows, start=2):
            row = {
                key: values[index] if index < len(values) else None
                for key, index in column_map.items()
            }
            if all(value is None or str(value).strip() == "" for value in row.values()):
                continue
            result.append((row_number, row))
        workbook.close()
        return result

    @staticmethod
    def _parse_float(value):
        if value is None or str(value).strip() == "":
            return None
        if isinstance(value, (int, float)):
            return float(value)
        return float(str(value).replace(",", "").strip())

    @staticmethod
    def _parse_date(value):
        if value is None or str(value).strip() == "":
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, (int, float)):
            return from_excel(value).date()
        text = str(value).strip().split(" ")[0]
        for date_format in DATE_FORMATS:
            try:
                return datetime.strptime(text, date_format).date()
            except ValueError:
                continue
        raise ValueError(text)

    def _build_lookups(self):
        employees_by_name = {}
        employees = self.env["hr.employee"].search([("company_id", "in", self.env.companies.ids)])
        for employee in employees:
            employees_by_name.setdefault(_normalize(employee.name), self.env["hr.employee"])
            employees_by_name[_normalize(employee.name)] |= employee

        types_by_name = {}
        input_types = self._get_input_types()
        for input_type in input_types:
            for name in (
                input_type.name,
                input_type.with_context(lang="en_US").name,
                input_type.code,
            ):
                if name:
                    types_by_name.setdefault(_normalize(name), input_type)

        durations = {}
        for key, label in DURATION_LABELS.items():
            durations[_normalize(key)] = key
            durations[_normalize(label)] = key
        duration_field = self.env["hr.salary.attachment"]._fields["duration_type"]
        for key, label in duration_field._description_selection(self.env):
            durations[_normalize(label)] = key
        return employees_by_name, types_by_name, durations

    def _prepare_line(self, row_number, row, lookups):
        employees_by_name, types_by_name, durations = lookups
        errors = []
        vals = {
            "row_number": row_number,
            "employee_name": str(row.get("employee") or "").strip(),
            "input_type_name": str(row.get("input_type") or "").strip(),
            "note": str(row.get("note") or "").strip(),
        }

        if not vals["employee_name"]:
            errors.append(_("Employee is missing."))
        else:
            matches = employees_by_name.get(_normalize(vals["employee_name"]))
            if not matches:
                errors.append(_("Employee '%s' not found.", vals["employee_name"]))
            elif len(matches) > 1:
                errors.append(_("More than one employee is named '%s'.", vals["employee_name"]))
            else:
                vals["employee_id"] = matches.id

        if not vals["input_type_name"]:
            errors.append(_("Type is missing."))
        else:
            input_type = types_by_name.get(_normalize(vals["input_type_name"]))
            if input_type:
                vals["input_type_id"] = input_type.id
            else:
                errors.append(_("Type '%s' not found or not available in adjustments.", vals["input_type_name"]))

        try:
            monthly_amount = self._parse_float(row.get("monthly_amount"))
            if monthly_amount is None:
                errors.append(_("Payslip Amount is missing."))
            elif monthly_amount <= 0:
                errors.append(_("Payslip Amount must be greater than 0."))
            else:
                vals["monthly_amount"] = monthly_amount
        except ValueError:
            errors.append(_("Payslip Amount '%s' is not a number.", row.get("monthly_amount")))

        raw_duration = _normalize(row.get("duration"))
        duration = durations.get(raw_duration) if raw_duration else "one"
        if duration:
            vals["duration_type"] = duration
        else:
            errors.append(_("Duration '%s' is invalid (use One Time, Limited or Unlimited).", row.get("duration")))

        if duration == "limited":
            try:
                total_amount = self._parse_float(row.get("total_amount"))
                if total_amount is None:
                    errors.append(_("Total Amount is required when Duration is Limited."))
                elif total_amount < (vals.get("monthly_amount") or 0) or total_amount <= 0:
                    errors.append(_("Total Amount must be greater than or equal to the Payslip Amount."))
                else:
                    vals["total_amount"] = total_amount
            except ValueError:
                errors.append(_("Total Amount '%s' is not a number.", row.get("total_amount")))

        try:
            date_start = self._parse_date(row.get("date_start"))
            if date_start:
                vals["date_start"] = date_start
            else:
                errors.append(_("Start Date is missing."))
        except ValueError:
            errors.append(_("Start Date '%s' is not a valid date (use YYYY-MM-DD).", row.get("date_start")))

        vals["status"] = "error" if errors else "ok"
        vals["message"] = "\n".join(errors)
        return vals

    def _build_preview(self):
        self.ensure_one()
        rows = self._read_rows()
        if not rows:
            raise UserError(_("The Excel file does not contain any rows to import."))
        lookups = self._build_lookups()
        self.line_ids.unlink()
        self.write({
            "state": "preview",
            "line_ids": [fields.Command.create(self._prepare_line(row_number, row, lookups)) for row_number, row in rows],
        })

    # ------------------------------------------------------------------
    # Actions
    # ------------------------------------------------------------------

    def action_check_file(self):
        self._build_preview()
        return self._reopen()

    def action_back(self):
        self.line_ids.unlink()
        self.state = "upload"
        return self._reopen()

    def action_import(self):
        self._build_preview()
        if self.error_count:
            return self._reopen()

        vals_list = []
        for line in self.line_ids:
            vals = {
                "employee_ids": [fields.Command.set(line.employee_id.ids)],
                "company_id": line.employee_id.company_id.id or self.env.company.id,
                "other_input_type_id": line.input_type_id.id,
                "monthly_amount": line.monthly_amount,
                "duration_type": line.duration_type,
                "date_start": line.date_start,
                "description": line.note or False,
            }
            if line.duration_type == "limited":
                vals["total_amount"] = line.total_amount
            vals_list.append(vals)

        adjustments = self.env["hr.salary.attachment"].create(vals_list)
        return {
            "type": "ir.actions.act_window",
            "name": _("Imported Salary Adjustments"),
            "res_model": "hr.salary.attachment",
            "view_mode": "list,form",
            "domain": [("id", "in", adjustments.ids)],
            "target": "current",
        }


class SalaryAdjustmentImportLine(models.TransientModel):
    _name = "salary.adjustment.import.line"
    _description = "Salary Adjustments Import Preview Line"
    _order = "status, row_number"

    wizard_id = fields.Many2one("salary.adjustment.import.wizard", required=True, ondelete="cascade")
    row_number = fields.Integer(string="Row")
    employee_name = fields.Char(string="Employee (file)")
    employee_id = fields.Many2one("hr.employee", string="Employee")
    input_type_name = fields.Char(string="Type (file)")
    input_type_id = fields.Many2one("hr.payslip.input.type", string="Type")
    monthly_amount = fields.Float(string="Payslip Amount", digits=(16, 3))
    duration_type = fields.Selection(list(DURATION_LABELS.items()), string="Duration")
    total_amount = fields.Float(string="Total Amount", digits=(16, 3))
    date_start = fields.Date(string="Start Date")
    note = fields.Char(string="Note")
    status = fields.Selection([("error", "Error"), ("ok", "OK")], required=True, default="ok")
    message = fields.Text(string="Errors")
